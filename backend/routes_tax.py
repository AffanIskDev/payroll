import csv
import io
import uuid
from datetime import datetime
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import JSONResponse, Response

from tax_forms import list_1721a1, render_1721a1

import bank_batch
from core import db, now_iso, require_manager, get_configs, csv_delimiter
from periods import payroll_month, run_type_of
from pipeline import get_period

# Pajak (saldo awal YTD, 1721-A1) and the BSI transfer export are Manager HRD only (user decision 9 Oct 2026)
router = APIRouter(prefix="/api/hr", dependencies=[Depends(require_manager)])

YTD_COLUMNS = ["nik", "tax_year", "cutoff_month", "ptkp_status", "gross_income_ytd", "employee_jht_jp_ytd",
               "biaya_jabatan_ytd", "pph21_paid_ytd", "notes"]
NUM_COLUMNS = ("gross_income_ytd", "employee_jht_jp_ytd", "biaya_jabatan_ytd", "pph21_paid_ytd")
BSI_COLUMNS = ["Rekening Debet", "Rekening Kredit", "Nama Penerima", "Nominal", "Berita Acara", "Email"]


def _rows_from_upload(name: str, raw: bytes) -> list:
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h or "").strip().lower() for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]
    text = raw.decode("utf-8-sig")
    first = text.splitlines()[0]
    delim = max([",", ";", "\t", "|"], key=first.count)
    return [{k.strip().lower(): v for k, v in r.items()} for r in csv.DictReader(io.StringIO(text), delimiter=delim)]


def _num(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    return float(s)


@router.get("/tax-balances")
async def list_balances(year: int):
    rows = await db.employee_tax_initial_balances.find({"tax_year": year}, {"_id": 0}).to_list(5000)
    emps = {e["id"]: e async for e in db.employees.find({}, {"_id": 0, "id": 1, "nik": 1, "name": 1})}
    for r in rows:
        e = emps.get(r["employee_id"], {})
        r.update(nik=e.get("nik"), name=e.get("name"))
    return sorted(rows, key=lambda r: r.get("nik") or "")


@router.get("/tax-balances/template")
async def template():
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=csv_delimiter(await get_configs()))
    w.writerow(YTD_COLUMNS)
    w.writerow(["EMP-001", 2026, 6, "K/1", 49874400, 1080000, 2493720, 498744, "contoh: saldo Jan-Jun dari sistem lama"])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": 'attachment; filename="template_saldo_awal_ytd.csv"'})


@router.post("/tax-balances/import")
async def import_balances(file: UploadFile = File(...)):
    try:
        rows = _rows_from_upload(file.filename or "", await file.read())
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"File tidak dapat dibaca: {e}")
    emps = {e["nik"]: e["id"] async for e in db.employees.find({}, {"_id": 0, "id": 1, "nik": 1})}
    imported, errors = 0, []
    for i, r in enumerate(rows, start=2):
        nik = str(r.get("nik") or "").strip()
        if nik not in emps:
            errors.append(f"Baris {i}: NIK '{nik}' tidak ditemukan")
            continue
        try:
            year, cutoff = int(_num(r.get("tax_year"))), int(_num(r.get("cutoff_month")))
            if not 1 <= cutoff <= 11:
                raise ValueError("cutoff_month harus 1-11")
            doc = {"employee_id": emps[nik], "tax_year": year, "cutoff_month": cutoff,
                   "ptkp_status": str(r.get("ptkp_status") or "").strip(), "notes": str(r.get("notes") or ""),
                   **{k: _num(r.get(k)) for k in NUM_COLUMNS}, "updated_at": now_iso()}
        except ValueError as e:
            errors.append(f"Baris {i}: {e}")
            continue
        await db.employee_tax_initial_balances.update_one(
            {"employee_id": doc["employee_id"], "tax_year": year},
            {"$set": doc, "$setOnInsert": {"id": str(uuid.uuid4())}}, upsert=True)
        imported += 1
    return {"imported": imported, "errors": errors}


@router.delete("/tax-balances/{bid}")
async def delete_balance(bid: str):
    res = await db.employee_tax_initial_balances.delete_one({"id": bid})
    if not res.deleted_count:
        raise HTTPException(404, "Data tidak ditemukan")
    return {"ok": True}


@router.get("/tax-forms/1721a1")
async def forms_1721a1(year: int):
    return await list_1721a1(year)


@router.get("/tax-forms/1721a1/{eid}/pdf")
async def form_1721a1_pdf(eid: str, year: int):
    forms = await list_1721a1(year, eid)
    if not forms:
        raise HTTPException(404, "Belum ada data pajak karyawan pada tahun ini")
    f = forms[0]
    return Response(await render_1721a1(f), media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="1721-A1_{f["nik"]}_{year}.pdf"'})


@router.post("/periods/{pid}/export-bsi")
async def export_bsi(pid: str):
    period = await get_period(pid)
    cfg = await get_configs()
    debit = cfg.get("BSI_COMPANY_DEBIT_ACCOUNT", "")
    slips = await db.payslips.find({"payroll_period_id": pid, "status": "approved"}, {"_id": 0}) \
        .sort("employee.nik", 1).to_list(5000)
    if not slips:
        raise HTTPException(400, "Belum ada slip berstatus approved pada periode ini")
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=csv_delimiter(cfg))
    w.writerow(BSI_COLUMNS)
    skipped = []
    for s in slips:
        e = s["employee"]
        amount = int(round(s["take_home_pay"]))
        if amount <= 0 or not e.get("bank_account_number"):
            skipped.append(e["nik"])
            continue
        w.writerow([debit, e["bank_account_number"], e["name"], amount,
                    f"Gaji {period['period_name']} {e['nik']}"[:50], e.get("email") or ""])
    fname = f"BSI_Payroll_{period['period_name'].replace(' ', '_')}.csv"
    return Response(buf.getvalue(), media_type="text/csv", headers={
        "Content-Disposition": f'attachment; filename="{fname}"',
        "X-Exported-Count": str(len(slips) - len(skipped)), "X-Skipped-Niks": ",".join(skipped)})


@router.post("/periods/{pid}/export-batch-bsi")
async def export_batch_bsi(pid: str):
    """Batch Payment BSI TXT (bank_batch), the file of ROSY's "Export Batch Payment". Amounts come from the period's
    slips, bank data from the employee master as last synced from ROSY (as ROSY reads mt_employee at export), so an
    account fixed in ROSY HRD reaches the file through "Sinkron dari ROSY" — also once the slips are approved & locked.

    Draft slips are exported too, as in ROSY since 30 Sep 2026 (user 10 Oct 2026: "samakan dengan rosy dulu ... biar
    bisa di test"). Restricting the file to approved & locked periods waits for the user's word."""
    period = await get_period(pid)
    # no length cap: a truncated list would silently leave people out of the transfer
    slips = await db.payslips.find({"payroll_period_id": pid}, {
        "_id": 0, "employee_id": 1, "employee": 1, "take_home_pay": 1, "payroll_status": 1}).to_list(None)
    if not slips:
        raise HTTPException(400, "Belum ada slip pada periode ini")
    emps = {e["id"]: e async for e in db.employees.find(
        {"id": {"$in": [s["employee_id"] for s in slips]}},
        {"_id": 0, "id": 1, "name": 1, "bank_code": 1, "bank_name": 1, "bank_account_number": 1, "bank_account_name": 1})}
    month = payroll_month(period["start_date"])
    # (pesan & awalan ID unggah, nama berkas): THR / Bonus Tahunan dibayar di tanggalnya sendiri, jangan tertukar dengan gaji
    kind = {"thr": ("THR", "THR"), "bonus_tahunan": ("BONUS", "Bonus Tahunan")}.get(run_type_of(period), ("GAJI", ""))
    try:
        res = bank_batch.build(slips, emps, month, datetime.now(ZoneInfo("Asia/Jakarta")).date().isoformat(), kind[0])
    except ValueError as e:
        raise HTTPException(422, str(e))
    if res["problems"]:
        return JSONResponse(status_code=422, content={
            "detail": f"Berkas tidak dibuat: {len(res['problems'])} data harus dibetulkan dulu",
            "problems": res["problems"], "slip_count": len(slips)})
    if not res["count"]:
        raise HTTPException(400, "Tidak ada yang ditransfer: seluruh slip periode ini Rp 0")
    return Response(res["text"], media_type="text/plain", headers={
        "Content-Disposition": f'attachment; filename="{bank_batch.file_name(period.get("company"), month, kind[1])}"',
        "X-Exported-Count": str(res["count"]), "X-Total-Amount": str(res["total"]),
        "X-Skipped-Niks": ",".join(res["skipped"]), "Cache-Control": "no-store"})
